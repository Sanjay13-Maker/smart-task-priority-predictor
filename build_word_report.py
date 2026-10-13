from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path('output/docx/Smart_Task_Priority_Predictor_Project_Report_SANJAY_S.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)
BLUE='17365D'; LIGHT='F4F7FA'; BORDER='D9D9D9'
doc=Document()
sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)

normal=doc.styles['Normal']; normal.font.name='Arial'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.2
for name, size in [('Title',24),('Heading 1',17),('Heading 2',12)]:
    s=doc.styles[name]; s.font.name='Arial'; s._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial'); s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def borders(cell):
    tcPr=cell._tc.get_or_add_tcPr(); tcBorders=tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None: tcBorders=OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; el=tcBorders.find(qn(tag))
        if el is None: el=OxmlElement(tag); tcBorders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),BORDER)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,BLUE); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v); shade(cells[i],LIGHT); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in t.rows:
        for c in row.cells:
            borders(c)
            for p in c.paragraphs: p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def para(text, bold=False, center=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    r=p.add_run(text); r.bold=bold; return p
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def newpage(): doc.add_page_break()

# Header/footer
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Smart Task Priority Predictor - Project Report | SANJAY S').font.size=Pt(8)

# Page 1
doc.add_paragraph('\n\n')
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('SMART TASK PRIORITY PREDICTOR')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Project Report').font.size=Pt(14)
doc.add_paragraph()
para('A Flask-based task management system with rule-based scoring and basic machine learning classification. The application helps users identify tasks requiring immediate attention by considering deadlines, urgency, business impact, effort, and dependencies.')
table(['Submission Detail','Information'],[['Submitted by','SANJAY S'],['Email ID','sanjayss007700@gmail.com'],['Intern ID','SHA20260724WB7B566C24INTERN'],['College','JCT COLLEGE OF ENGINEERING AND TECHNOLOGY']],[2,4.7])
doc.add_heading('Abstract',1)
para('Task prioritization is a practical problem for students, project teams, and individual users. Smart Task Priority Predictor provides a clear and repeatable way to create tasks, evaluate their priority, and track the result. The system combines an understandable weighted scoring formula with a Decision Tree classifier. It presents Low, Medium, High, or Critical priority together with a score and simple reasons. The project demonstrates Python, Flask, SQLite, Pandas, Scikit-learn, HTML, CSS, and Pytest in one intermediate-level application.')
para('Keywords: task management, task priority, Flask, SQLite, Decision Tree, machine learning, CSV batch prediction.')
newpage()
# Page 2
doc.add_heading('1. Introduction and Problem Statement',0)
para('Users commonly manage tasks with different deadlines, effort requirements, urgency levels, business impacts, and dependencies. Comparing all of these factors manually is time-consuming and can lead to inconsistent decisions. A normal task list records work but may not give a user a consistent method for deciding what should be completed first.')
para('Smart Task Priority Predictor addresses this challenge through a web application that captures relevant task characteristics and transforms them into an explainable priority recommendation. The application supports human judgement rather than replacing it. It provides a structured baseline so that every task is reviewed against the same criteria.')
doc.add_heading('1.1 Purpose and Objectives',1)
para('The project develops a Flask web application for task management and priority prediction. A registered user can create, view, edit, delete, search, filter, and update their own tasks. For each valid task, the application calculates a priority score from 0 to 100 and uses a saved Decision Tree model to predict a priority category.')
bullets(['Build a usable application for personal and project task management.','Generate a transparent priority score and one of four priority labels.','Train and integrate a basic Decision Tree model using prepared task data.','Support CSV batch processing, manual overrides, and automated testing.'])
doc.add_heading('1.2 Scope',1)
table(['Included','Excluded'],[['Flask login, task CRUD, SQLite storage','Cloud deployment, containers, CI/CD'],['Weighted scoring and Decision Tree classification','OAuth/JWT and enterprise databases'],['CSV batch prediction and dashboard summaries','External task-management integrations'],['Unit and integration tests','Deep learning, LLMs and drift automation']],[3.3,3.3])
newpage()
# Page 3
doc.add_heading('2. System Architecture and Workflow',0)
para('The system follows a simple three-layer design. The presentation layer contains HTML and CSS pages for registration, login, dashboard, task forms, task lists, and batch prediction. The Flask application layer handles routes, sessions, validation, task operations, scoring, and prediction requests. SQLite forms the data layer, storing users, tasks, predictions, and priority overrides. A machine-learning module loads a saved Scikit-learn model when prediction is requested.')
table(['Component','Responsibility'],[['HTML, CSS and Bootstrap','Forms, task lists, dashboard, validation messages and result display.'],['Flask application','Routes, validation, sessions, CRUD operations and prediction flow.'],['Priority rule module','Score calculation, score mapping and user-readable reasons.'],['Pandas module','CSV reading, cleaning, feature preparation and batch processing.'],['Scikit-learn model','Classification as Low, Medium, High or Critical.'],['SQLite database','Persistent storage of users, tasks, predictions and overrides.']],[2,4.6])
doc.add_heading('2.1 Prediction Flow',1)
para('The user enters task information. Flask validates all values and derives days remaining. The priority rule module calculates a score, the model input is prepared in the same format used during training, and the Decision Tree predicts a class. If model output disagrees with the score range, the score range is used as the final explainable category. The final result and reasons are saved in SQLite.')
newpage()
# Page 4
doc.add_heading('3. Functional Requirements',0)
para('The requirements cover user management, task management, prediction, overrides, and batch processing. Each operation is designed for a local application while preserving correct ownership of data.')
table(['ID','Requirement'],[['FR-01 to FR-04','Register, login, logout, Flask session handling, and access only to the user\'s own tasks.'],['FR-05 to FR-11','Create, view, edit, delete, search, filter, sort, and update task status.'],['FR-12 to FR-17','Calculate score, classify priority, show results and reasons, and store each prediction.'],['FR-18','Allow manual priority override with a mandatory reason.'],['FR-19','Upload CSV, validate rows, process valid tasks, preview and download results.']],[1.4,5.2])
doc.add_heading('3.1 Input Validation and Error Handling',1)
para('Task name and deadline are mandatory. Estimated effort must be greater than zero. Business impact and urgency must be integers from 1 to 10. Dependency count must be zero or greater. The deadline cannot be earlier than the creation date. Completed tasks are not given new priority and cancelled tasks are excluded from ranking.')
table(['Situation','System Response'],[['Missing value','Show a validation message beside the field.'],['Invalid impact or urgency','Explain the permitted range is 1 to 10.'],['Invalid effort or deadline','Request a positive effort value or valid deadline.'],['Incorrect login','Show Invalid email or password.'],['Invalid CSV','List missing or incorrect columns.'],['Service failure','Show general error and log technical details.']],[2.2,4.4])
newpage()
# Page 5
doc.add_heading('4. Priority Calculation Logic',0)
para('The scoring logic is intentionally simple and explainable. Business impact has the largest weight, followed by urgency. Dependencies and deadline pressure add points, while effort is a small adjustment.')
doc.add_heading('4.1 Score Formula',1)
para('Priority Score = (Business Impact x 4) + (Urgency x 3) + (Dependency Count x 2) + Deadline Points + Effort Adjustment',True)
para('The score is limited to the range 0 to 100. An example task with impact 9, urgency 9, four dependencies, and a deadline within two days receives 36 + 27 + 8 + 15 = 86 before effort adjustment and is classified as Critical.')
table(['Deadline Situation','Points','Meaning'],[['Overdue','20','Immediate attention is required.'],['Today or within 2 days','15','Task is time-sensitive.'],['Within 7 days','10','Deadline is approaching.'],['Within 14 days','7','Moderate time pressure.'],['More than 14 days','3','Lower deadline pressure.']],[2,0.8,3.8])
table(['Condition','Result'],[['Effort below 3 hours','Add 5 points for a potential quick win.'],['Effort above 20 hours','Subtract 5 points.'],['Score 80 to 100','Critical'],['Score 60 to 79','High'],['Score 40 to 59','Medium'],['Score 0 to 39','Low']],[2.6,4])
newpage()
# Page 6
doc.add_heading('5. Machine Learning and Data Preparation',0)
para('The machine learning component is a supervised classification problem using a Decision Tree Classifier. The model is selected because it is easy to train, visualize, and explain. It predicts Low, Medium, High, or Critical priority and complements the transparent rule-based score.')
table(['Dataset Field','Purpose'],[['task_id','Unique identifier and duplicate detection.'],['created_date and deadline_date','Derive days to deadline.'],['estimated_effort','Expected work in hours.'],['business_impact and urgency','Core ratings from 1 to 10.'],['dependency_count','Number of dependent tasks.'],['task_type','Categorical feature converted to numeric values.'],['priority_label','Target class for training.']],[2.2,4.4])
doc.add_heading('5.1 Preprocessing and Training',1)
para('Pandas reads the CSV file, checks required columns, converts dates, validates numerical ranges, fills suitable missing values, removes duplicate identifiers, creates days-to-deadline, and one-hot encodes task type. The data is split into 80 percent training and 20 percent testing data. The trained model is evaluated using accuracy, precision, recall, and F1-score, then saved using Joblib. The target is at least 75 percent accuracy as a practical internship goal.')
newpage()
# Page 7
doc.add_heading('6. Database Design and Security',0)
para('SQLite is used as a lightweight relational database for the local version. The schema keeps users, their tasks, generated predictions, and manual changes connected through clear relationships.')
table(['Table','Key Information'],[['Users','user_id, name, unique email, password value, created_at'],['Tasks','task_id, user_id, title, deadline, effort, impact, urgency, dependencies, type and status'],['Predictions','prediction_id, task_id, score, priority label, model prediction and created_at'],['Override History','override_id, task_id, original priority, new priority, reason and created_at']],[1.65,4.95])
doc.add_heading('6.1 Security and Logging',1)
bullets(['Flask sessions maintain login state after authentication.','Server-side validation checks every submitted value.','Parameterized SQLite queries prevent unsafe SQL construction.','Internal exception details are written to logs rather than shown to users.','The secret key is stored in configuration or an environment variable.','Users can only view and modify their own tasks.'])
para('Python logging records application startup, login success or failure, task operations, prediction generation, CSV processing, and unexpected errors. Advanced security platforms, OAuth, JWT, and API gateways remain outside the intended internship scope.')
newpage()
# Page 8
doc.add_heading('7. User Interface and Batch Prediction',0)
para('The interface presents a clear workflow for users who may not know machine learning. It uses familiar pages for login, task entry, task lists, dashboard summaries, result explanations, and CSV upload.')
table(['Page','User Outcome'],[['Login and Register','Create account and access the application with validation feedback.'],['Dashboard','Review total tasks, priority counts, overdue tasks, recent tasks and chart.'],['Task List','Search, filter, sort, edit and delete owned tasks.'],['Task Form','Enter fields, receive result, reasons and optional override.'],['Batch Prediction','Upload CSV, view validation messages, preview and download result.']],[1.7,4.9])
doc.add_heading('7.1 Batch Prediction and Override',1)
para('For batch prediction, the user uploads a CSV file with required columns. The application validates each row and applies the same scoring and prediction logic used for a single task. Invalid rows are reported clearly, while valid rows receive a score and priority label. Results can be downloaded as a CSV file. The initial version supports up to 500 records per upload.')
para('The manual override feature preserves human control. A user can select another priority after viewing a recommendation, but must provide a reason. The original prediction, selected priority, reason, and timestamp are stored as an audit trail.')
newpage()
# Page 9
doc.add_heading('8. Testing, Risks and Validation',0)
para('Pytest supports unit tests for isolated logic and integration tests for complete workflows. The test suite checks normal actions and invalid input so that changes to scoring, validation, database operations, and routes can be verified consistently.')
table(['Testing Area','Example Checks'],[['Priority Score','Low, medium, high, overdue, quick-win and dependency cases.'],['Validation','Invalid impact, urgency, effort, deadline and missing fields.'],['Feature Preparation','Days-to-deadline and task type encoding.'],['Prediction','Valid output is Low, Medium, High or Critical.'],['Database CRUD','Create, read, update and delete operations.'],['Integration','Login -> task -> prediction -> save; CSV -> prediction -> download.']],[1.7,4.9])
doc.add_heading('8.1 Risks and Mitigation',1)
table(['Risk','Mitigation'],[['Small or imbalanced dataset','Create realistic examples and review class counts.'],['Poor model accuracy','Clean data, verify features and tune Decision Tree settings.'],['Incorrect input','Validate all fields before prediction.'],['Wrong suggestion','Show reasons and allow a manual override.'],['Invalid CSV or missing model','Validate columns and show a clear service error.']],[2.3,4.3])
para('The project is accepted when users can register, manage tasks, receive valid scores and labels, save predictions and overrides, process CSV files, view dashboard counts, and run core automated tests.')
newpage()
# Page 10
doc.add_heading('9. Implementation Plan, Conclusion and Future Work',0)
para('Implementation begins with Flask setup and basic pages, followed by SQLite tables and task CRUD operations. Validation and scoring are added before preparing the dataset and training the Decision Tree. The saved model is connected to Flask, then dashboard charts, CSV processing, tests, documentation, and local demonstration are completed.')
table(['Phase','Outcome'],[['1 to 3','Project setup, database tables and task CRUD.'],['4 to 6','Validation, dataset preparation and scoring.'],['7 to 8','Decision Tree training and model integration.'],['9 to 10','Dashboard and CSV batch prediction.'],['11 to 12','Tests, defect fixes, documentation and demonstration.']],[1.5,5.1])
doc.add_heading('9.1 Technology Stack',1)
table(['Layer','Technology'],[['Application','Python and Flask'],['Data and ML','Pandas, NumPy, Scikit-learn and Joblib'],['Storage and Interface','SQLite, HTML, CSS, Bootstrap and Matplotlib'],['Testing','Pytest']],[2.4,4.2])
doc.add_heading('9.2 Conclusion',1)
para('Smart Task Priority Predictor turns task information into an understandable priority recommendation. It combines task management, a transparent score, basic machine learning, readable explanations, and manual user control. The project demonstrates Flask routing, SQLite CRUD operations, data preprocessing, classification, error handling, testing, and documentation in one coherent local application.')
doc.add_heading('9.3 Future Enhancements',1)
para('Future versions may use a larger real-world dataset, compare a Random Forest model, add project-level grouping and filters, introduce notifications, collect feedback, and deploy to simple hosting. Advanced AI, automatic retraining, monitoring, and external integrations remain outside the scope of this version.')
doc.core_properties.author='SANJAY S'; doc.core_properties.title='Smart Task Priority Predictor Project Report'
doc.save(OUT)
print(OUT.resolve())
